"""The local server: static files, MIME types, token / Host / Origin checks, settings validation, jobs."""
import http.client
import json
import sys
import threading
from http.server import ThreadingHTTPServer
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import app  # noqa: E402


@pytest.fixture(scope='module')
def server(tmp_path_factory):
    tmp = tmp_path_factory.mktemp('appdir')
    saved = (app.SETTINGS_FILE, app.HISTORY_FILE, app.PORT_FILE, app.DB_FILE)
    app.SETTINGS_FILE, app.HISTORY_FILE, app.PORT_FILE = tmp / 'settings.json', tmp / 'history.json', tmp / 'port.txt'
    app.DB_FILE = tmp / 'market.db'
    (tmp / 'settings.json').write_text(json.dumps({'pages': 9, 'cache_hours': 'abc', 'theme': 'dark', 'old_key': 1}))
    app.S = app.State(start_loop=False)
    httpd = ThreadingHTTPServer(('127.0.0.1', 0), app.Handler)
    app.S.port = httpd.server_address[1]
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    yield app.S
    httpd.shutdown()
    app.S.db.close()
    app.SETTINGS_FILE, app.HISTORY_FILE, app.PORT_FILE, app.DB_FILE = saved


def call(S, method, path, body=None, token=True, host=None, origin=None, headers=None):
    c = http.client.HTTPConnection('127.0.0.1', S.port, timeout=5)
    h = {'Host': host or '127.0.0.1:%d' % S.port}
    if token:
        h['X-PC-Token'] = S.token
    if origin:
        h['Origin'] = origin
    if body is not None:
        body = json.dumps(body).encode()
        h['Content-Type'] = 'application/json'
    h.update(headers or {})
    c.request(method, path, body=body, headers=h)
    r = c.getresponse()
    data = r.read()
    c.close()
    return r.status, r.getheader('Content-Type') or '', data


def test_old_settings_are_migrated_and_clamped(server):
    assert server.settings['pages'] == 3
    assert server.settings['cache_hours'] == 6
    assert server.settings['theme'] == 'dark'
    assert 'old_key' not in server.settings
    assert server.settings['locations']['AU'] == '2000'


def test_index_has_the_token_and_assets_have_fixed_types(server):
    st, ct, body = call(server, 'GET', '/', token=False)
    assert st == 200 and ct.startswith('text/html')
    assert server.token.encode() in body and b'%%PC_TOKEN%%' not in body
    for path, kind in (('/ui/js/app.js', 'text/javascript'), ('/ui/js/legacy.js', 'text/javascript'),
                       ('/ui/css/tokens.css', 'text/css'), ('/ui/css/app.css', 'text/css'), ('/favicon.svg', 'image/svg+xml')):
        st, ct, _ = call(server, 'GET', path, token=False)
        assert st == 200 and ct.startswith(kind), path


def test_every_script_and_stylesheet_in_index_exists():
    import re
    html = (app.UI_DIR / 'index.html').read_text(encoding='utf-8')
    refs = re.findall(r'(?:src|href)="/ui/([^"]+)"', html)
    assert refs
    for ref in refs:
        assert (app.UI_DIR / ref).is_file(), ref


def test_static_files_cannot_escape_the_ui_folder(server):
    for path in ('/ui/../app.py', '/ui/%2e%2e/app.py', '/ui/..%2fapp.py', '/ui/js/../../config.py', '/app.py', '/ui/'):
        st, _, body = call(server, 'GET', path, token=False)
        assert st == 404, path
        assert b'import' not in body


def test_api_needs_the_token(server):
    assert call(server, 'GET', '/api/state', token=False)[0] == 403
    st, _, body = call(server, 'GET', '/api/state')
    assert st == 200 and json.loads(body)['version'] == app.VERSION
    assert call(server, 'GET', '/api/state', headers={'X-PC-Token': 'wrong'}, token=False)[0] == 403
    assert call(server, 'POST', '/api/quit', body={}, token=False)[0] == 403


def test_other_hosts_and_origins_are_refused(server):
    assert call(server, 'GET', '/', host='evil.example:%d' % server.port, token=False)[0] == 403
    assert call(server, 'GET', '/api/state', host='evil.example')[0] == 403
    assert call(server, 'POST', '/api/settings', body={'pages': 2}, origin='https://evil.example')[0] == 403
    assert call(server, 'POST', '/api/settings', body={'pages': 2}, origin='http://127.0.0.1:%d' % server.port)[0] == 200
    assert call(server, 'GET', '/api/state', host='localhost:%d' % server.port)[0] == 200


def test_ping_works_without_the_token(server):
    st, _, body = call(server, 'GET', '/api/ping', token=False)
    assert st == 200 and json.loads(body)['ok'] is True


def test_settings_are_validated(server):
    st, _, body = call(server, 'POST', '/api/settings', body={'theme': 'pink'})
    assert st == 400 and b'theme' in body
    st, _, body = call(server, 'POST', '/api/settings', body={'fast_days': 99, 'theme': 'system', 'locations': {'AU': ' 3000 '}})
    s = json.loads(body)['settings']
    assert st == 200 and s['fast_days'] == 10 and s['theme'] == 'system' and s['locations']['AU'] == '3000'
    assert json.loads(app.SETTINGS_FILE.read_text())['theme'] == 'system'
    st, _, body = call(server, 'POST', '/api/settings', body={'fast_days': 'soon'})
    assert st == 400
    call(server, 'POST', '/api/settings', body={'locations': {'AU': '2000'}, 'fast_days': 3})


def test_unknown_routes_jobs_and_actions(server):
    assert call(server, 'GET', '/api/nope')[0] == 404
    assert call(server, 'GET', '/api/job?id=missing')[0] == 410
    assert call(server, 'POST', '/api/browser', body={'action': 'explode'})[0] == 400
    assert call(server, 'POST', '/api/export', body={'kind': 'check', 'job': 'missing'})[0] == 410
    assert call(server, 'POST', '/api/open', body={'url': 'http://example.com'})[0] == 400
    assert call(server, 'POST', '/api/check', body={'keyword': '   '})[0] == 400


def test_job_manager_prunes_old_and_extra_jobs():
    import concurrent.futures as cf
    import jobs

    def run(coro):
        coro.close()
        f = cf.Future()
        f.set_result({'ok': 1})
        return f

    async def work(progress):
        return 1
    jm = jobs.JobManager(run)
    ids = [jm.start('t', lambda p: work(p)) for _ in range(jobs.MAX_JOBS + 5)]
    assert jm.get(ids[-1])['status'] == 'done'
    jm.prune()
    assert len(jm.jobs) <= jobs.MAX_JOBS and jm.get(ids[-1])
    jm.prune(now=jm.get(ids[-1])['finished'] + jobs.KEEP_SECONDS + 1)
    assert not jm.jobs


def test_failed_job_reports_its_error():
    import concurrent.futures as cf
    import jobs
    said = []

    def run(coro):
        coro.close()
        f = cf.Future()
        f.set_exception(RuntimeError('Amazon asked for a captcha\nmore'))
        return f

    async def work(progress):
        return 1
    jm = jobs.JobManager(run, say=said.append)
    j = jm.get(jm.start('check', lambda p: work(p)))
    assert j['status'] == 'error' and 'captcha' in j['error'] and said and 'Check failed' in said[0]


def test_radar_scan_reads_the_body_once_and_answers_at_once(server, monkeypatch):
    """The market used to be read from the body twice: the second read waited for the socket to time out, then the
    scan failed with market None ("Starting…" for minutes)."""
    import asyncio
    import time as _t
    got = []
    monkeypatch.setattr(server, 'new_job', lambda kind, fn: got.append(fn) or 'job1')
    t0 = _t.time()
    st, _, raw = call(server, 'POST', '/api/radar/scan', {'market': 'AU'})
    assert st == 200 and json.loads(raw)['job'] == 'job1' and _t.time() - t0 < 2
    lines = []

    async def lists(task, progress=None):
        lines.append(task['market'])
        raise RuntimeError('stop here')
    monkeypatch.setattr(server.collectors, 'lists', lists)
    with pytest.raises(RuntimeError, match='stop here'):
        asyncio.run(got[0](lines.append))
    assert 'Amazon Australia' in lines[0] and lines[1] == 'AU'


def test_state_has_exchange_rates_and_your_ranges(server):
    st, _, raw = call(server, 'GET', '/api/state')
    fx = json.loads(raw)['fx']
    assert st == 200 and fx['currency'] == 'USD' and fx['ranges']['US'] == [20, 50]
    assert fx['ranges']['AE'][0] > 70 and fx['per_usd']['AED'] == 3.6725
    st, _, raw = call(server, 'POST', '/api/settings', {'price_currency': 'EUR', 'price_min': 10})
    d = json.loads(raw)
    assert st == 200 and d['settings']['price_min'] == 10 and d['fx']['ranges']['US'][0] > 10
    call(server, 'POST', '/api/settings', {'price_currency': 'USD', 'price_min': 20})


def test_opportunity_map_route(server):
    st, _, raw = call(server, 'GET', '/api/opportunity_map?market=AE')
    d = json.loads(raw)
    assert st == 200 and d['market'] == 'AE' and isinstance(d['points'], list) and d['local_line'] == 4.5


def test_profit_takes_the_supplier_cost_in_any_currency(server):
    import fxrates
    base = {'market': 'AU', 'price': 49.99, 'l': 30, 'w': 20, 'h': 10, 'kg': 0.5}
    st, _, raw = call(server, 'POST', '/api/profit', {**base, 'cogs_usd': 5})
    usd = json.loads(raw)
    aed = fxrates.convert(5, 'USD', 'AED')
    st2, _, raw = call(server, 'POST', '/api/profit', {**base, 'cogs': aed, 'cogs_currency': 'AED'})
    d = json.loads(raw)
    assert st == st2 == 200 and d['economics']['COGS'] == pytest.approx(usd['economics']['COGS'], abs=0.01)
    assert d['cost_currency'] == 'AED'
    assert d['max_cogs_20_cost'] == pytest.approx(fxrates.convert(usd['max_cogs_20'], 'USD', 'AED'), abs=0.02)
    st, _, raw = call(server, 'POST', '/api/profit', {**base, 'cogs': 5, 'cogs_currency': 'XYZ'})
    assert json.loads(raw)['cost_currency'] == 'USD'                       # unknown currency: read as US$


def test_check_export_has_prices_in_your_currency(server):
    import io
    import time

    import pandas as pd

    import amazon_check as ac
    import fxrates
    rows = [{'market': 'AU', 'asin': 'B%d' % i, 'title': 't', 'price': p, 'sponsored': False, 'days': 2, 'delivery': 'x',
             'reviews': 1, 'bought': None} for i, p in enumerate((10.0, 30.0))]
    s = ac.summarize('AU', 'moving bags', 'Sydney 2000', rows)
    server.jobs.jobs['xjob'] = {'id': 'xjob', 'status': 'done', 'result': {'keyword': 'moving bags', 'results': [s]},
                                'finished': time.time(), 'log': []}
    st, _, raw = call(server, 'POST', '/api/export', {'kind': 'check', 'job': 'xjob'})
    sheets = pd.read_excel(io.BytesIO(raw), sheet_name=None)
    summ, listings = sheets['Summary'], sheets['Listings']
    assert st == 200 and summ.loc[0, 'currency'] == 'A$' and summ.loc[0, 'typical_price'] == 20
    assert summ.loc[0, 'typical_price_USD'] == pytest.approx(fxrates.convert(20, 'AUD', 'USD'), abs=0.01)
    assert list(listings['currency']) == ['A$', 'A$']
    assert listings.loc[1, 'price_USD'] == pytest.approx(fxrates.convert(30, 'AUD', 'USD'), abs=0.01)


def test_list_export_checks_every_row_and_converts_prices(server):
    import io

    import openpyxl
    import pandas as pd

    import fxrates
    rows = [{'product': '=HYPERLINK("http://x")', 'market': 'AU', 'local': 3, 'overseas': 9, 'fast': 2, 'total': 20,
             'demand': 1200.5, 'per_local_seller': 300, 'overseas_share': 0.42, 'price_med': 20, 'verdict': 'gap'},
            {'product': 'dog bed', 'market': 'AE', 'error': 'captcha'},
            {'product': 'bad', 'market': 'XX', 'local': 1}, 'not a row', {'product': 'odd', 'market': 'US', 'local': 'many'}]
    st, _, raw = call(server, 'POST', '/api/export', {'kind': 'batch', 'rows': rows})
    df = pd.read_excel(io.BytesIO(raw))
    assert st == 200 and list(df['Market']) == ['Australia', 'UAE', 'USA']                  # the unknown market is left out
    assert df.loc[0, 'Overseas share of sales %'] == 42 and df.loc[1, 'Verdict'] == 'captcha'
    assert df.loc[0, 'Typical price (US$)'] == pytest.approx(fxrates.convert(20, 'AUD', 'USD'), abs=0.01)
    assert pd.isna(df.loc[2, 'Local sellers'])                                              # not a number: left empty
    ws = openpyxl.load_workbook(io.BytesIO(raw)).active
    assert ws['A2'].data_type == 's' and ws['A2'].value.startswith("'=")                     # text, never a formula
    assert call(server, 'POST', '/api/export', {'kind': 'batch', 'rows': []})[0] == 400


def test_download_name_works_for_an_arabic_keyword(server):
    import time

    import amazon_check as ac
    rows = [{'market': 'AE', 'asin': 'B1', 'title': 't', 'price': 50.0, 'sponsored': False, 'days': 2, 'delivery': 'x',
             'reviews': 1, 'bought': None}]
    kw = 'حقيبة "سفر"'
    server.jobs.jobs['ajob'] = {'id': 'ajob', 'status': 'done', 'finished': time.time(), 'log': [],
                                'result': {'keyword': kw, 'results': [ac.summarize('AE', kw, 'Dubai', rows)]}}
    c = http.client.HTTPConnection('127.0.0.1', server.port, timeout=5)
    c.request('POST', '/api/export', body=json.dumps({'kind': 'check', 'job': 'ajob'}).encode(),
              headers={'Host': '127.0.0.1:%d' % server.port, 'X-PC-Token': server.token, 'Content-Type': 'application/json'})
    r = c.getresponse()
    r.read()
    c.close()
    cd = r.getheader('Content-Disposition')
    assert r.status == 200 and 'filename="Delivery_check_.xlsx"' in cd and "filename*=UTF-8''Delivery_check_%D8" in cd
    assert app.attachment('Top_products.xlsx') == 'attachment; filename="Top_products.xlsx"; filename*=UTF-8\'\'Top_products.xlsx'


def test_the_top_listings_and_their_sellers_are_read(server, monkeypatch):
    import asyncio

    import amazon_check as ac
    import product
    import storage
    rows = [{'market': 'AU', 'asin': 'S%d' % i, 'title': 'Box %d' % i, 'price': 30.0, 'sponsored': i == 1, 'days': 2,
             'delivery': 'x', 'reviews': 5, 'bought': None} for i in range(12)]
    server.db.write(storage.save_serp, ac.summarize('AU', 'storage box', 'Sydney 2000', rows), 'user', None, 1, '2000')
    server.review_wall.clear()
    calls = []
    who = {'S0': 'SELLERA1', 'S3': 'SELLERA1', 'S4': 'SELLERB1'}

    async def page(checker, asin, code, reviews=False):
        calls.append(('page', asin))
        if asin == 'S2':
            raise RuntimeError('page did not load')
        amazon = asin == 'S5'
        return {'title': 'Box', 'root_rank': 100, 'status': 'ok', 'pkg_cm': [40.0, 30.0, 12.0], 'pkg_kg': 1.2,
                'buybox_seller': 'Amazon AU' if amazon else None, 'seller_id': None if amazon else who.get(asin, 'SELLERC1'),
                'fulfillment': 'amazon' if amazon else 'fba', 'brand': 'Boxy'}

    async def critical(checker, asin, code):
        calls.append(('reviews', asin))
        raise product.SignInNeeded('sign in')

    async def seller(checker, sid, code):
        calls.append(('seller', sid))
        return {'name': 'Store ' + sid[-2], 'country': 'AU' if sid == 'SELLERB1' else 'CN', 'status': 'ok'}
    monkeypatch.setattr(product, 'read_product', page)
    monkeypatch.setattr(product, 'read_critical_reviews', critical)
    monkeypatch.setattr(product, 'read_seller', seller)
    assert asyncio.run(server.read_top('AU', 'storage box')) == 12                  # 9 pages + 3 sellers
    top = ['S0'] + ['S%d' % i for i in range(2, 11)]                                # the ad (S1) is never opened
    assert calls == [('page', a) for a in top] + [('reviews', 'S0')] + [('seller', s) for s in ('SELLERA1', 'SELLERB1', 'SELLERC1')]
    st, _, raw = call(server, 'GET', '/api/report/sellers?market=AU&kw=storage%20box')
    w = json.loads(raw)
    assert st == 200 and w['takeaway'] == 'Mostly Chinese sellers using FBA: 7 of the top 10. Amazon itself sells 1.'
    assert w['todo'] == {'pages': ['S2'], 'sellers': []} and w['sellers'][0] == {
        'name': 'Store C', 'id': 'SELLERC1', 'country': 'CN', 'ships': ['fba'], 'listings': 5, 'units': None, 'share': 0.5}
    assert w['listings'][4]['seller'] == 'Amazon' and w['listings'][0]['seller_url'].endswith('/sp?seller=SELLERA1')
    assert server.db.read(lambda c: __import__('reports').size_suggestion(c, 'AU', 'storage box'))['dims_cm'] == [40.0, 30.0, 12.0]
    r = server.explain('storage box', [{'market': 'AU'}, {'market': 'US', 'error': 'blocked'}])
    assert 'in Australia' in r[0]['sentence'] and 'mostly Chinese FBA sellers' in r[0]['sentence'] and 'sentence' not in r[1]
    assert r[0]['who'] == {'takeaway': w['takeaway'], 'known': 9, 'n': 10}

    async def captcha(checker, asin, code, reviews=False):
        calls.append(('page', asin))
        raise ac.CaptchaRequired('Amazon wants a captcha')
    monkeypatch.setattr(product, 'read_product', captcha)
    del calls[:]
    with pytest.raises(ac.CaptchaRequired):                                        # nothing read: the reason is shown
        asyncio.run(server.read_top('AU', 'storage box', force=True))
    assert calls == [('page', 'S0')]                                                # stops at the captcha
    assert call(server, 'POST', '/api/report/top', {'market': 'AU', 'kw': ' '})[0] == 400
    st, _, raw = call(server, 'POST', '/api/parse_size', {'text': 'Single package size: 30X20X10 cm Single gross weight: 0.5 kg'})
    assert st == 200 and json.loads(raw) == {'dims_cm': [30, 20, 10], 'kg': 0.5}


def test_the_report_is_saved_as_a_pdf(server, monkeypatch, tmp_path):
    import asyncio
    import time
    seen = {}

    class Page:
        async def goto(self, url):
            seen['url'] = url

        async def wait_for_function(self, js, timeout=None):
            seen['wait'] = js

        async def emulate_media(self, media):
            seen['media'] = media

        async def pdf(self, **kw):
            seen['pdf'] = kw
            return b'%PDF-1.4 fake'

    class Browser:
        async def new_page(self, viewport):
            seen['viewport'] = viewport
            return Page()

        async def close(self):
            seen['closed'] = True

    class Chromium:
        async def launch(self, headless, **extra):
            if extra.get('channel') == 'chrome':
                raise RuntimeError('no Chrome here')
            seen['launch'] = (headless, extra)
            return Browser()

    class Playwright:
        chromium = Chromium()
    monkeypatch.setattr(server.checker, 'pw', Playwright())
    monkeypatch.setattr(app, 'downloads_dir', lambda: tmp_path)
    r = asyncio.run(server.make_pdf('AU', 'dog bed: big/XL', ['summary', 'sellers'], None))
    assert r['name'] == 'dog bed big XL - Amazon Australia - %s.pdf' % time.strftime('%Y-%m-%d') and r['folder'] == str(tmp_path)
    assert (tmp_path / r['name']).read_bytes() == b'%PDF-1.4 fake' and seen['closed'] and seen['media'] == 'print'
    assert seen['url'] == 'http://127.0.0.1:%d/?print=summary,sellers#/k/AU/dog%%20bed%%3A%%20big%%2FXL' % server.port
    assert seen['launch'] == (True, {'channel': 'msedge'}) and seen['viewport']['width'] == app.A4_CONTENT_PX   # Edge when no Chrome
    assert seen['pdf']['format'] == 'A4' and 'Amazon Australia' in seen['pdf']['footer_template']
    assert asyncio.run(server.make_pdf('AU', 'dog bed: big/XL', ['summary'], None))['name'].endswith(' (2).pdf')
    assert call(server, 'POST', '/api/pdf/open', {'id': 'nope'})[0] == 404                       # only PDFs the app made
    assert call(server, 'POST', '/api/report/pdf', {'market': 'AU', 'kw': 'dog bed', 'sections': ['x']})[0] == 400
    assert call(server, 'POST', '/api/report/pdf', {'market': 'AU', 'kw': 'dog bed', 'sections': 'summary'})[0] == 400
    assert call(server, 'POST', '/api/report/pdf', {'market': 'AU', 'kw': ' ', 'sections': ['summary']})[0] == 400
    assert app.pdf_name('a/b: "c"?', 'AE', '2026-10-12') == 'a b c - Amazon UAE - 2026-10-12.pdf'


def test_negative_reviews_are_read_from_amazon(server, monkeypatch):
    import asyncio

    import amazon_check as ac
    import product
    import storage
    rows = [{'market': 'US', 'asin': 'M%d' % i, 'title': 'Bowl %d' % i, 'price': 20.0, 'sponsored': i == 1, 'days': 2,
             'delivery': 'x', 'reviews': 5, 'bought': None} for i in range(4)]
    server.db.write(storage.save_serp, ac.summarize('US', 'mixing bowl', 'New York 10001', rows), 'user', None, 1, '10001')
    calls, signed = [], [False]

    async def critical(checker, asin, code):
        calls.append(('reviews page', asin))
        if not signed[0]:
            raise product.SignInNeeded('sign in')
        return [{'id': 'RX' + asin + '000000', 'stars': 1.0, 'title': 'Broke', 'body': 'b', 'local': True, 'helpful': 9}]

    async def page(checker, asin, code, reviews=False):
        calls.append(('product page', asin))
        return {'title': 'Bowl', 'root_rank': 10, 'status': 'ok', 'reviews': 50, 'signed_in': signed[0], 'review_items': [
            {'id': 'R' + asin + '00000000', 'stars': 2.0, 'title': 'Cracked', 'body': 'b', 'local': True, 'helpful': 3}]}
    monkeypatch.setattr(product, 'read_critical_reviews', critical)
    monkeypatch.setattr(product, 'read_product', page)
    st, _, raw = call(server, 'GET', '/api/report/reviews?market=US&kw=mixing%20bowl')
    assert st == 200 and json.loads(raw)['todo'] == ['M0', 'M2', 'M3'] and json.loads(raw)['items'] == []
    assert asyncio.run(server.read_reviews('US', ['M0', 'M2', 'M3'])) == 3
    assert calls == [('reviews page', 'M0'), ('product page', 'M0'), ('product page', 'M2'), ('product page', 'M3')]
    r = json.loads(call(server, 'GET', '/api/report/reviews?market=US&kw=mixing%20bowl')[2])
    assert r['todo'] == [] and [i['asin'] for i in r['items']] == ['M0', 'M2', 'M3'] and not r['critical']
    signed[0], calls[:] = True, []                          # signed in now: a product page shows it, then reviews pages
    asyncio.run(server.read_reviews('US', ['M0', 'M2']))
    assert calls == [('product page', 'M0'), ('reviews page', 'M2')]
    r = json.loads(call(server, 'GET', '/api/report/reviews?market=US&asin=M2')[2])
    assert r['critical'] and r['items'][0]['title'] == 'Broke'

    async def captcha(checker, asin, code, reviews=False):
        raise ac.CaptchaRequired('Amazon wants a captcha')
    monkeypatch.setattr(product, 'read_critical_reviews', captcha)
    with pytest.raises(ac.CaptchaRequired):                                     # nothing read: the reason is shown
        asyncio.run(server.read_reviews('US', ['M3']))
    assert call(server, 'POST', '/api/report/reviews', {'market': 'US', 'asin': 'B0-1'})[0] == 400
    assert call(server, 'GET', '/api/report/reviews?market=US&kw=%20')[0] == 400


def test_qualify_checks_its_input_and_takes_at_most_12(server, monkeypatch):
    import asyncio
    got, calls = [], []
    monkeypatch.setattr(server, 'new_job', lambda kind, fn: got.append((kind, fn)) or 'qjob')

    async def qualify(market, asins, progress=None, checks=0, pages=12, phrases=None):
        calls.append((market, asins, checks, phrases))
        return {'status': 'ok', 'note': '12 products: 3 trends', 'counts': {}}
    monkeypatch.setattr(server.collectors, 'qualify_products', qualify)
    assert call(server, 'POST', '/api/discovery/qualify', {'market': 'AU', 'asins': []})[0] == 400
    assert call(server, 'POST', '/api/discovery/qualify', {'market': 'AU', 'asins': 'B0ABCDEFGH'})[0] == 400
    assert call(server, 'POST', '/api/discovery/qualify', {'market': 'AU', 'asins': ['B0-1', 'x' * 11]})[0] == 400
    asins = ['b0abcdefgh', 'B0ABCDEFGH', 'bad', None, 5] + ['B0%08d' % i for i in range(20)]
    body = {'market': 'AU', 'asins': asins, 'phrases': {'B0ABCDEFGH': '  Tower  FAN ', 'bad': 'x', 'B000000001': 7}}
    st, _, raw = call(server, 'POST', '/api/discovery/qualify', body)
    assert st == 200 and json.loads(raw)['job'] == 'qjob' and got[0][0] == 'qualify'
    out = asyncio.run(got[0][1](lambda m: None))
    market, todo, checks, phrases = calls[0]
    assert market == 'AU' and len(todo) == 12 and todo[0] == 'B0ABCDEFGH' and checks == 0   # never searches other markets
    assert phrases == {'B0ABCDEFGH': 'tower fan'} and out['skipped'] == 9 and out['note'] == '12 products: 3 trends'


def test_your_choices_per_product_are_checked_and_saved(server):
    pref = lambda s, asin='B0ABCDEFGH', mk='US': call(server, 'POST', '/api/pref', {'market': mk, 'asin': asin, 'set': s})  # noqa: E731
    assert pref({'liked': True}, asin='bad')[0] == 400 and pref({'liked': True}, mk='XX')[0] == 400
    assert pref({})[0] == 400 and pref({'cost': 'abc'})[0] == 400 and pref({'cost': -2})[0] == 400
    st, _, raw = pref({'cost': 36.725, 'cost_currency': 'AED', 'phrase': '  Tower  FAN ', 'liked': 1}, asin='b0abcdefgh')
    p = json.loads(raw)['pref']
    assert st == 200 and p['asin'] == 'B0ABCDEFGH' and p['cogs_usd'] == pytest.approx(10.0, abs=0.01)
    assert p['phrase'] == 'tower fan' and p['liked_at']
    p = json.loads(pref({'cost': '', 'phrase': '', 'liked': False})[2])['pref']
    assert p['cogs_usd'] is None and p['phrase'] is None and p['liked_at'] is None


def test_find_the_underserved_version_checks_its_input(server, monkeypatch):
    import asyncio
    got, calls = [], []
    monkeypatch.setattr(server, 'new_job', lambda kind, fn: got.append((kind, fn)) or 'vjob')

    async def underserved(market, phrase, boom=None, progress=None, n=4):
        calls.append((market, phrase, boom))
        return {'phrase': phrase, 'variants': []}
    monkeypatch.setattr(server.collectors, 'underserved', underserved)
    assert call(server, 'POST', '/api/discovery/variants', {'market': 'AU', 'phrase': '  '})[0] == 400
    body = {'market': 'AU', 'phrase': ' Tower  FAN ', 'boom': {'trend_score': 250, 'cluster_surge': 'x'}}
    st, _, raw = call(server, 'POST', '/api/discovery/variants', body)
    assert st == 200 and json.loads(raw)['job'] == 'vjob' and got[0][0] == 'variants'
    asyncio.run(got[0][1](lambda m: None))
    assert calls == [('AU', 'tower fan', {'trend_score': 100.0, 'cluster_surge': None})]


def test_where_to_launch_checks_its_input(server):
    assert call(server, 'POST', '/api/launch', {'market': 'US', 'kw': ' '})[0] == 400
    st, _, raw = call(server, 'POST', '/api/launch', {'market': 'US', 'kw': 'zzz nothing here', 'cogs': 'x', 'units': -5, 'l': 10})
    r = json.loads(raw)
    assert st == 200 and r['units'] == 1 and r['cost_usd'] is None and r['markets']['US'] == {'searched': False}


def test_your_ideas_scan_joins_the_list_history_and_is_scored(server, monkeypatch):
    import asyncio

    import ideas
    cats = [{'slug': 'kitchen', 'name': 'Kitchen', 'risky': False}]
    items = [{'asin': 'B0K%07d' % i, 'title': 'Bamboo Utensil Holder %s' % 'abcdefghijklmnopqrstuvwxyzabcdefghi'[i],
              'image': '', 'market': 'AU', 'url': '', 'category': 'Kitchen', 'kind': 'movers', 'slug': 'kitchen', 'deal': False,
              'rank': i + 1, 'pct': 100, 'rank_now': 500 + i, 'rank_before': 1000, 'rating': 4.5, 'reviews': 12, 'price': 30.0}
             for i in range(35)]
    scans = []

    async def categories(checker, code):
        return cats

    async def scan(checker, code, chosen, kinds, hide, progress):
        scans.append(code)
        return {'market': code, 'scanned': len(items), 'lists': 1, 'categories': chosen, 'items': [dict(i) for i in items]}
    monkeypatch.setattr(ideas, 'categories', categories)
    monkeypatch.setattr(ideas, 'scan', scan)
    monkeypatch.setitem(server.settings, 'price_filter', False)
    server.idea_cache.clear()
    server.idea_cats.clear()
    count = "SELECT COUNT(*) AS n FROM list_snapshot WHERE market='AU' AND slug='kitchen'"
    before = server.db.one(count)['n']
    out = asyncio.run(server.find_ideas('AU', [], ['movers'], True, False, lambda m: None))
    asyncio.run(server.find_ideas('AU', [], ['movers'], True, False, lambda m: None))     # the same scan, from the cache
    assert scans == ['AU'] and server.db.one(count)['n'] == before + 1
    snap = server.db.one("SELECT status, n FROM list_snapshot WHERE market='AU' AND slug='kitchen' ORDER BY id DESC LIMIT 1")
    assert snap['status'] == 'ok' and snap['n'] == 35
    json.dumps(app.clean(out), allow_nan=False)
    idea = out['ideas'][0]
    assert idea['opportunity']['score'] is not None and idea['presence'][-1] == 1 and idea['score'] is not None
    assert next(p for p in idea['opportunity']['parts'] if p['key'] == 'A')['src'] == 'lists'
    assert out['coverage']['n'] == len(out['ideas']) and len(out['presence_days']) == 14
    st, _, raw = call(server, 'GET', '/api/radar?market=AU')                         # the radar sees your scan too
    r = json.loads(raw)
    shown = {i['asin'] for i in r['items']} | {m['asin'] for i in r['items'] for m in i['members']}
    assert st == 200 and idea['asin'] in shown and r['coverage']['n'] == len(r['items'])
    assert idea['similar'] == 34 and len(idea['members']) == 34                     # 35 utensil holders: one idea


def test_the_board_routes_check_their_input_and_never_show_the_password(server, monkeypatch, tmp_path):
    import mailer
    import winsys
    monkeypatch.setattr(app, 'MAIL_SECRET', tmp_path / 'mail.secret')
    monkeypatch.setattr(winsys, 'protect', lambda b: b'X' + b[::-1])
    monkeypatch.setattr(winsys, 'unprotect', lambda b: b[1:][::-1])
    sent = []
    monkeypatch.setattr(mailer, '_smtp', lambda acct, password, msg: sent.append((password, msg)))
    mailer._sent.clear()
    assert call(server, 'GET', '/api/board', token=False)[0] == 403
    st, _, raw = call(server, 'POST', '/api/board/add', {'title': 'Dog Bed', 'market': 'US', 'kw': 'dog bed'})
    cid = json.loads(raw)['id']
    b = json.loads(call(server, 'GET', '/api/board')[2])
    assert st == 200 and b['cards'][0]['id'] == cid and b['mail']['ready'] is False
    st, _, raw = call(server, 'POST', '/api/board/line', {'id': cid, 'n': 1, 'market': 'US', 'set': {'ship_mode': 'boat'}})
    assert st == 400 and 'sea or air' in json.loads(raw)['error']
    assert call(server, 'POST', '/api/board/line', {'id': cid, 'n': 1, 'market': 'US', 'set': {'units': -5}})[0] == 400
    assert call(server, 'POST', '/api/board/move', {'id': cid, 'stage': 'reorder'})[0] == 400
    assert call(server, 'GET', '/api/board/card?id=999')[0] == 404
    ppl = json.loads(call(server, 'POST', '/api/board/notifiers', {'add': {'name': 'Sara', 'email': 'sara@example.com'}})[2])
    assert call(server, 'POST', '/api/board/notifiers', {'add': {'name': 'x', 'email': 'nope'}})[0] == 400
    task = {'id': cid, 'stage': 'sourcing', 'days': 7, 'notify': [ppl['notifiers'][0]['id']]}
    st, _, raw = call(server, 'POST', '/api/board/task', task)
    assert st == 400 and 'Gmail' in json.loads(raw)['error']
    st, _, raw = call(server, 'POST', '/api/mail', {'password': 'abcd efgh ijkl mnop'})
    assert st == 200 and json.loads(raw)['has_password'] is True and b'abcd' not in raw
    assert b'abcd' not in call(server, 'GET', '/api/state')[2] and b'"board_rev"' in call(server, 'GET', '/api/state')[2]
    assert call(server, 'POST', '/api/settings', {'mail_address': 'not an email'})[0] == 400
    call(server, 'POST', '/api/settings', {'mail_address': 'me@gmail.com', 'mail_from_name': 'Ali'})
    try:
        st, _, raw = call(server, 'POST', '/api/board/task', task)
        t = json.loads(raw)
        assert st == 200 and t['mails'] == 3 and t['due_day']
        mails = json.loads(call(server, 'GET', '/api/board/preview?task=%d' % t['task_id'])[2])['mails']
        assert [m['kind'] for m in mails] == ['start', 'reminder', 'meeting']
        assert mails[0]['subject'].startswith('New task: Sourcing & PO for Dog Bed') and 'Hi Sara,' in mails[0]['text']
        st, _, raw = call(server, 'POST', '/api/mail/test', {})
        assert st == 200 and json.loads(raw)['to'] == 'me@gmail.com'
        assert sent[0][0] == 'abcdefghijklmnop' and sent[0][1]['To'] == 'Ali <me@gmail.com>'
        summary = json.loads(call(server, 'GET', '/api/board/summary')[2])['text']
        assert 'Dog Bed (USA)' in summary and 'due ' in summary
        st, _, raw = call(server, 'GET', '/api/mail')
        assert json.loads(raw)['pending'] == 3 and json.loads(raw)['ready'] is True
        assert call(server, 'POST', '/api/mail', {'forget': True})[0] == 200 and not (tmp_path / 'mail.secret').exists()
    finally:
        call(server, 'POST', '/api/settings', {'mail_address': '', 'mail_from_name': ''})
        call(server, 'POST', '/api/board/archive', {'id': cid, 'archived': True})


def test_notes_counts_and_products_can_be_removed(server):
    cid = json.loads(call(server, 'POST', '/api/board/add', {'title': 'Lick Mat', 'market': 'US', 'kw': 'lick mat'})[2])['id']
    assert call(server, 'POST', '/api/board/note', {'id': cid, 'text': 'Ask for samples'})[0] == 200
    card = json.loads(call(server, 'GET', '/api/board/card?id=%d' % cid)[2])
    note = next(e['id'] for e in card['events'] if e['kind'] == 'note')
    assert call(server, 'POST', '/api/board/note', {'remove': note})[0] == 200
    assert call(server, 'POST', '/api/board/note', {'remove': note})[0] == 400
    assert call(server, 'POST', '/api/board/stock', {'remove': 99999})[0] == 400
    assert call(server, 'POST', '/api/board/archive', {'id': cid, 'archived': True})[0] == 200
    assert any(a['id'] == cid for a in json.loads(call(server, 'GET', '/api/board')[2])['archived'])
    assert call(server, 'POST', '/api/board/delete', {'id': cid})[0] == 200
    assert call(server, 'GET', '/api/board/card?id=%d' % cid)[0] == 404
    assert all(a['id'] != cid for a in json.loads(call(server, 'GET', '/api/board')[2])['archived'])
    assert call(server, 'POST', '/api/board/delete', {'id': 'x'})[0] == 400


def test_inform_attendees_previews_and_sends_the_product_update(server, monkeypatch, tmp_path):
    import mailer
    import winsys
    monkeypatch.setattr(app, 'MAIL_SECRET', tmp_path / 'mail.secret')
    monkeypatch.setattr(winsys, 'protect', lambda b: b'X' + b[::-1])
    monkeypatch.setattr(winsys, 'unprotect', lambda b: b[1:][::-1])
    sent = []
    monkeypatch.setattr(mailer, '_smtp', lambda acct, password, msg: sent.append(msg))
    mailer._sent.clear()
    cid = json.loads(call(server, 'POST', '/api/board/add', {'title': 'Yoga Mat', 'market': 'AU', 'kw': 'yoga mat'})[2])['id']
    st, _, raw = call(server, 'POST', '/api/board/card', {'id': cid, 'set': {'alibaba': 'javascript:alert(1)'}})
    assert st == 400 and 'https://' in json.loads(raw)['error']
    people = json.loads(call(server, 'POST', '/api/board/notifiers', {'add': {'name': 'Li Wei', 'email': 'li@example.com'}})[2])
    li = next(p['id'] for p in people['notifiers'] if p['email'] == 'li@example.com')
    view = json.loads(call(server, 'GET', '/api/board/inform?to=%d&message=Hello' % li)[2])
    assert view['to'] == ['Li Wei'] and view['subject'].startswith('Product update, ') and 'Yoga Mat' in view['text']
    assert 'Hello' in view['text'] and view['mail']['ready'] is False and '<html>' in view['html']
    assert json.loads(call(server, 'GET', '/api/board/inform?to=0')[2])['to'] == []
    assert call(server, 'POST', '/api/board/inform', {'to': [li]})[0] == 400                # no Gmail yet
    call(server, 'POST', '/api/mail', {'password': 'abcd efgh ijkl mnop'})
    call(server, 'POST', '/api/settings', {'mail_address': 'me@gmail.com', 'mail_from_name': 'Ali'})
    try:
        st, _, raw = call(server, 'POST', '/api/board/inform', {'to': [li], 'message': 'Quotes by Friday'})
        assert st == 200 and json.loads(raw)['to'] == ['Li Wei'] and sent[-1]['To'] == 'Li Wei <li@example.com>'
        body = sent[-1].get_body(('plain',)).get_content()
        assert 'Yoga Mat' in body and 'Quotes by Friday' in body and 'Search on Alibaba' in body
        st, _, raw = call(server, 'POST', '/api/board/inform', {'to': []})
        assert st == 400 and 'Pick who' in json.loads(raw)['error']
        assert json.loads(call(server, 'GET', '/api/board')[2])['inform']['to'] == ['Li Wei']
    finally:
        call(server, 'POST', '/api/settings', {'mail_address': '', 'mail_from_name': ''})
        call(server, 'POST', '/api/mail', {'forget': True})
        call(server, 'POST', '/api/board/archive', {'id': cid, 'archived': True})
