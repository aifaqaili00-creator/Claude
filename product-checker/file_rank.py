"""Rank a Helium 10 export (Black Box, Xray or similar CSV/Excel) and return the top products.

The rules match the AU & UAE Product Finder: enough monthly sales, few reviews, a price that
leaves room after fees, and no product types that bring extra rules (electrical, kids, chemicals...).
"""
import re
from pathlib import Path

import pandas as pd

# Targets per marketplace. price in local currency, sales = units per month, reviews = max reviews.
TARGETS = {
    'AU': {'currency': 'A$', 'price': (25, 70), 'sales': 100, 'reviews': 230},
    'AE': {'currency': 'AED', 'price': (60, 260), 'sales': 50, 'reviews': 115},
    'US': {'currency': 'US$', 'price': (20, 60), 'sales': 300, 'reviews': 300},
}
MARKET_NAMES = {'AU': 'Australia', 'AE': 'UAE', 'US': 'USA'}

# canonical column -> names used by Helium 10 exports (compared after normalising)
ALIASES = {
    'title': ['title', 'product details', 'product name', 'product title'],
    'asin': ['asin'],
    'url': ['url', 'product url', 'link'],
    'brand': ['brand'],
    'category': ['category'],
    'subcategory': ['subcategory'],
    'price': ['price'],
    'sales': ['monthly sales', 'asin sales', 'sales', 'parent level sales', 'est monthly sales'],
    'revenue': ['monthly revenue', 'asin revenue', 'revenue', 'parent level revenue'],
    'reviews': ['review count', 'reviews', 'ratings count', 'number of reviews'],
    'rating': ['reviews rating', 'ratings', 'rating', 'review rating'],
    'sellers': ['number of active sellers', 'active sellers', 'sellers', 'number of sellers'],
    'age': ['listing age months', 'age months', 'listing age'],
    'created': ['creation date', 'date first available'],
    'weight': ['weight', 'weight lb'],
    'variations': ['variation count', 'variations'],
    'trend': ['sales trend 90 days', 'sales trend'],
    'seller': ['seller', 'buy box seller', 'bb seller'],
    'seller_country': ['seller country region', 'seller country'],
    'fulfillment': ['fulfillment', 'fulfilment'],
}

R = lambda p: re.compile(p, re.I)
CHECKS = [
    ('consumable/chemical', R(r'\b(food|treats?|snacks?|chews?|jerky|flea|tick|supplements?|vitamins?|tablets?|capsules?|gummies|creatine|protein|pre-?workout|medicine|nappies|wipes|shampoo|litter|cleaner|detergent|sanitis\w*|sanitiz\w*|spray|solution|disposable|candles?|diffuser|paint|varnish|markers?|filters?|incense|remover|liquid|oil|wax|polish|alcohol|isopropyl|resin|epoxy|glue|refills?|soap|toothpaste|cream|gel|powder|fertili[sz]er|pesticide|insecticide|repellent|bait|poison|seeds?|soil|coffee|lotion|serum|deodor\w*|fragrance|perfume|boric)\b')),
    ('electrical', R(r'\b(heater|heated|heating|electric\w*|rechargeable|cordless|usb|battery|batteries|charger|plug|led|lamp|lights?|bluetooth|smart|wi-?fi|motor|blender|kettle|fryer|vacuum|steamer|stimulator|tens|ems|thermostat|alarm|power ?board|socket|tracker|warmer|humidifier|purifier|recorder|printer|faucet|torch|flashlight|fountain|solar|fan|sensor|automatic|massager|watts?|volt)\b')),
    ('kids/toy', R(r'\b(kids?|children|child|baby|toddler|girls|boys|youth|toys?|squish\w*)\b')),
    ('medical', R(r'\b(pain relief|posture|brace|orthopedic|orthotic|thermometer|blood|medical|first aid|bandage|syringes?|needles?|vials?|bacteriostatic|therapy|hearing|back support|compression|splint|pregnan\w*|maternity|cervical)\b')),
    ('fits another brand', R(r'(spare parts|compatible|replacement|for dyson|dreame|roborock|ecovacs|shark|kitchen ?aid|nespresso|thermomix|dolce gusto)')),
    ('seasonal', R(r'\b(christmas|xmas|advent|halloween|easter|valentine)\b')),
    ('bulky', R(r'\b(mattress|pillows?|duvet|quilt|throw|blanket|desk|chair mat|shelf|shelves|cabinet|table|ladder|wheelbarrow)\b')),
]


def _norm(s):
    s = re.sub(r'\(.*?\)', ' ', str(s).lower())          # drop "(months)", "(90 days) (%)"
    s = re.sub(r'(au\$|us\$|aed|\$|%)', ' ', s)
    return re.sub(r'[^a-z0-9]+', ' ', s).strip()


def _num(v):
    if v is None or (isinstance(v, float) and v != v):
        return float('nan')
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(',', '')
    m = re.search(r'-?\d+(\.\d+)?', s)
    if not m:
        return float('nan')
    n = float(m.group(0))
    if re.search(r'\d\s*k\b', s, re.I):
        n *= 1000
    return n


def map_columns(df):
    """Return {canonical: original column name} for the columns this file has."""
    normed = {c: _norm(c) for c in df.columns}
    out = {}
    for key, names in ALIASES.items():
        for name in names:                                    # exact match first, in alias order
            hit = next((c for c, n in normed.items() if n == name and c not in out.values()), None)
            if hit:
                out[key] = hit
                break
        if key not in out and key == 'price':                 # "Price AU$", "Price  $" ...
            hit = next((c for c, n in normed.items() if n.startswith('price') and 'trend' not in n and 'change' not in n), None)
            if hit:
                out[key] = hit
    return out


def detect_market(df, cols, filename=''):
    name = Path(filename).name.upper()
    for code, keys in (('AU', ('AU_', '_AU', 'AUSTRALIA')), ('AE', ('AE_', '_AE', 'UAE')), ('US', ('US_', '_US', 'USA'))):
        if any(k in name for k in keys):
            return code
    if 'url' in cols:
        urls = ' '.join(df[cols['url']].astype(str).head(20))
        if 'amazon.com.au' in urls:
            return 'AU'
        if 'amazon.ae' in urls:
            return 'AE'
        if 'amazon.com/' in urls:
            return 'US'
    return 'AU'


def load(path):
    path = str(path)
    if path.lower().endswith(('.xlsx', '.xlsm', '.xls')):
        return pd.read_excel(path)
    for enc in ('utf-8-sig', 'cp1252'):
        try:
            first = open(path, encoding=enc).readline()
            return pd.read_csv(path, encoding=enc, skiprows=1 if first.startswith('#') else 0)
        except UnicodeDecodeError:
            continue
    return pd.read_csv(path, encoding='latin-1')


def rank(path, market='auto', top=10):
    """Read the file and return (market, top_df, all_df, summary_text)."""
    raw = load(path)
    cols = map_columns(raw)
    if 'title' not in cols or 'sales' not in cols:
        raise ValueError('This file has no Title / Sales columns. Use a Helium 10 Black Box or Xray export.')
    mk = detect_market(raw, cols, path) if market == 'auto' else market
    T = TARGETS[mk]

    d = pd.DataFrame()
    for key, col in cols.items():
        d[key] = raw[col]
    for key in ('title', 'asin', 'url', 'brand', 'category', 'subcategory', 'seller', 'seller_country', 'fulfillment'):
        d[key] = d[key].fillna('').astype(str) if key in d else ''
    for key in ('price', 'sales', 'revenue', 'reviews', 'rating', 'sellers', 'age', 'weight', 'variations', 'trend'):
        d[key] = d[key].map(_num) if key in d else float('nan')
    if d['age'].isna().all() and 'created' in d:
        created = pd.to_datetime(d['created'], errors='coerce')
        d['age'] = ((pd.Timestamp.now() - created).dt.days / 30.4).round()
    d['revenue'] = d['revenue'].fillna(d['price'] * d['sales'])
    d = d[d['title'].str.strip() != '']

    def flags(r):
        f = [name for name, rx in CHECKS if rx.search(r.title)]
        if r.seller.lower().startswith('amazon') or r.fulfillment.upper() == 'AMZ':
            f.append('Amazon sells')
        if r.variations > 20:
            f.append('%d variations' % r.variations)
        if r.weight > 8:
            f.append('heavy')
        if r.price < T['price'][0]:
            f.append('cheap')
        if r.price > T['price'][1]:
            f.append('expensive')
        return ', '.join(f)

    d['flags'] = d.apply(flags, axis=1)
    rv = d['reviews'].fillna(0)
    new_bonus = d['age'].map(lambda a: 1.15 if 0 < a <= 12 else 1.0)
    fix_bonus = d['rating'].map(lambda s: 1.1 if 0 < s <= 4.2 else 1.0)
    trend = d['trend'].map(lambda t: 1.1 if t >= 30 else 0.85 if t <= -20 else 1.0)   # growing / shrinking sales
    d['score'] = (d['sales'].fillna(0) / (rv + 20) ** 0.5 * new_bonus * fix_bonus * trend).round(1)
    sweet = (d['sales'] >= T['sales']) & (rv <= T['reviews'])
    d['verdict'] = 'Skip'
    d.loc[(d['sales'] >= T['sales'] / 2) & (d['flags'] == ''), 'verdict'] = 'Close'
    d.loc[sweet, 'verdict'] = 'Check first'
    d.loc[sweet & (d['flags'] == ''), 'verdict'] = 'Good pick'
    order = {'Good pick': 0, 'Check first': 1, 'Close': 2, 'Skip': 3}
    d['_o'] = d['verdict'].map(order)
    d = d.sort_values(['_o', 'score'], ascending=[True, False]).drop(columns='_o').reset_index(drop=True)
    d.insert(0, 'rank', range(1, len(d) + 1))

    counts = d['verdict'].value_counts()
    summary = '%s (%s) - %d products: %d good picks, %d check first, %d close, %d skip.\nTargets: price %s%g-%g, at least %d sales/month, at most %d reviews.' % (
        MARKET_NAMES[mk], mk, len(d), counts.get('Good pick', 0), counts.get('Check first', 0), counts.get('Close', 0),
        counts.get('Skip', 0), T['currency'], T['price'][0], T['price'][1], T['sales'], T['reviews'])
    return mk, d.head(top), d, summary


EXPORT_COLS = ['rank', 'verdict', 'title', 'brand', 'price', 'sales', 'revenue', 'reviews', 'rating', 'age',
               'sellers', 'trend', 'weight', 'fulfillment', 'flags', 'score', 'asin', 'url']
EXPORT_HEAD = ['Rank', 'Verdict', 'Title', 'Brand', 'Price', 'Sales/mo', 'Revenue/mo', 'Reviews', 'Rating', 'Age (mo)',
               'Sellers', '90d trend %', 'Weight lb', 'Fulfillment', 'Flags', 'Score', 'ASIN', 'URL']


def export_excel(df, path):
    from openpyxl.styles import Font, PatternFill
    out = df[[c for c in EXPORT_COLS if c in df]].copy()
    out.columns = [EXPORT_HEAD[EXPORT_COLS.index(c)] for c in out.columns]
    with pd.ExcelWriter(path, engine='openpyxl') as xw:
        out.to_excel(xw, index=False, sheet_name='Ranked')
        ws = xw.sheets['Ranked']
        fills = {'Good pick': 'E3F4E3', 'Check first': 'FDF1D6', 'Close': 'ECEEED', 'Skip': 'F8E1E1'}
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='1F4FD6')
        for row in ws.iter_rows(min_row=2):
            color = fills.get(row[1].value)
            if color:
                row[1].fill = PatternFill('solid', fgColor=color)
        widths = {'Title': 60, 'Flags': 28, 'URL': 40}
        for i, head in enumerate(out.columns, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = widths.get(head, 12)
        ws.freeze_panes = 'D2'
