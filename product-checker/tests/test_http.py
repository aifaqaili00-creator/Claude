"""The local server: static files, MIME types, token / Host / Origin checks, settings validation, jobs."""
import http.client
import json
import sys
import threading
from http.server import ThreadingHTTPServer
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import app  # noqa: E402


@pytest.fixture(scope='module')
def server(tmp_path_factory):
    tmp = tmp_path_factory.mktemp('appdir')
    saved = (app.SETTINGS_FILE, app.HISTORY_FILE, app.PORT_FILE, app.DB_FILE)
    app.SETTINGS_FILE, app.HISTORY_FILE, app.PORT_FILE = tmp / 'settings.json', tmp / 'history.json', tmp / 'port.txt'
    app.DB_FILE = tmp / 'market.db'
    (tmp / 'settings.json').write_text(json.dumps({'pages': 9, 'cache_hours': 'abc', 'theme': 'dark', 'old_key': 1}))
    app.S = app.State(start_loop=False)
    httpd = ThreadingHTTPServer(('127.0.0.1', 0), app.Handler)
    app.S.port = httpd.server_address[1]
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    yield app.S
    httpd.shutdown()
    app.S.db.close()
    app.SETTINGS_FILE, app.HISTORY_FILE, app.PORT_FILE, app.DB_FILE = saved


def call(S, method, path, body=None, token=True, host=None, origin=None, headers=None):
    c = http.client.HTTPConnection('127.0.0.1', S.port, timeout=5)
    h = {'Host': host or '127.0.0.1:%d' % S.port}
    if token:
        h['X-PC-Token'] = S.token
    if origin:
        h['Origin'] = origin
    if body is not None:
        body = json.dumps(body).encode()
        h['Content-Type'] = 'application/json'
    h.update(headers or {})
    c.request(method, path, body=body, headers=h)
    r = c.getresponse()
    data = r.read()
    c.close()
    return r.status, r.getheader('Content-Type') or '', data


def test_old_settings_are_migrated_and_clamped(server):
    assert server.settings['pages'] == 3
    assert server.settings['cache_hours'] == 6
    assert server.settings['theme'] == 'dark'
    assert 'old_key' not in server.settings
    assert server.settings['locations']['AU'] == '2000'


def test_index_has_the_token_and_assets_have_fixed_types(server):
    st, ct, body = call(server, 'GET', '/', token=False)
    assert st == 200 and ct.startswith('text/html')
    assert server.token.encode() in body and b'%%PC_TOKEN%%' not in body
    for path, kind in (('/ui/js/app.js', 'text/javascript'), ('/ui/js/legacy.js', 'text/javascript'),
                       ('/ui/css/tokens.css', 'text/css'), ('/ui/css/app.css', 'text/css'), ('/favicon.svg', 'image/svg+xml')):
        st, ct, _ = call(server, 'GET', path, token=False)
        assert st == 200 and ct.startswith(kind), path


def test_every_script_and_stylesheet_in_index_exists():
    import re
    html = (app.UI_DIR / 'index.html').read_text(encoding='utf-8')
    refs = re.findall(r'(?:src|href)="/ui/([^"]+)"', html)
    assert refs
    for ref in refs:
        assert (app.UI_DIR / ref).is_file(), ref


def test_static_files_cannot_escape_the_ui_folder(server):
    for path in ('/ui/../app.py', '/ui/%2e%2e/app.py', '/ui/..%2fapp.py', '/ui/js/../../config.py', '/app.py', '/ui/'):
        st, _, body = call(server, 'GET', path, token=False)
        assert st == 404, path
        assert b'import' not in body


def test_api_needs_the_token(server):
    assert call(server, 'GET', '/api/state', token=False)[0] == 403
    st, _, body = call(server, 'GET', '/api/state')
    assert st == 200 and json.loads(body)['version'] == app.VERSION
    assert call(server, 'GET', '/api/state', headers={'X-PC-Token': 'wrong'}, token=False)[0] == 403
    assert call(server, 'POST', '/api/quit', body={}, token=False)[0] == 403


def test_other_hosts_and_origins_are_refused(server):
    assert call(server, 'GET', '/', host='evil.example:%d' % server.port, token=False)[0] == 403
    assert call(server, 'GET', '/api/state', host='evil.example')[0] == 403
    assert call(server, 'POST', '/api/settings', body={'pages': 2}, origin='https://evil.example')[0] == 403
    assert call(server, 'POST', '/api/settings', body={'pages': 2}, origin='http://127.0.0.1:%d' % server.port)[0] == 200
    assert call(server, 'GET', '/api/state', host='localhost:%d' % server.port)[0] == 200


def test_ping_works_without_the_token(server):
    st, _, body = call(server, 'GET', '/api/ping', token=False)
    assert st == 200 and json.loads(body)['ok'] is True


def test_settings_are_validated(server):
    st, _, body = call(server, 'POST', '/api/settings', body={'theme': 'pink'})
    assert st == 400 and b'theme' in body
    st, _, body = call(server, 'POST', '/api/settings', body={'fast_days': 99, 'theme': 'system', 'locations': {'AU': ' 3000 '}})
    s = json.loads(body)['settings']
    assert st == 200 and s['fast_days'] == 10 and s['theme'] == 'system' and s['locations']['AU'] == '3000'
    assert json.loads(app.SETTINGS_FILE.read_text())['theme'] == 'system'
    st, _, body = call(server, 'POST', '/api/settings', body={'fast_days': 'soon'})
    assert st == 400
    call(server, 'POST', '/api/settings', body={'locations': {'AU': '2000'}, 'fast_days': 3})


def test_unknown_routes_jobs_and_actions(server):
    assert call(server, 'GET', '/api/nope')[0] == 404
    assert call(server, 'GET', '/api/job?id=missing')[0] == 410
    assert call(server, 'POST', '/api/browser', body={'action': 'explode'})[0] == 400
    assert call(server, 'POST', '/api/export', body={'kind': 'check', 'job': 'missing'})[0] == 410
    assert call(server, 'POST', '/api/open', body={'url': 'http://example.com'})[0] == 400
    assert call(server, 'POST', '/api/check', body={'keyword': '   '})[0] == 400


def test_job_manager_prunes_old_and_extra_jobs():
    import concurrent.futures as cf
    import jobs

    def run(coro):
        coro.close()
        f = cf.Future()
        f.set_result({'ok': 1})
        return f

    async def work(progress):
        return 1
    jm = jobs.JobManager(run)
    ids = [jm.start('t', lambda p: work(p)) for _ in range(jobs.MAX_JOBS + 5)]
    assert jm.get(ids[-1])['status'] == 'done'
    jm.prune()
    assert len(jm.jobs) <= jobs.MAX_JOBS and jm.get(ids[-1])
    jm.prune(now=jm.get(ids[-1])['finished'] + jobs.KEEP_SECONDS + 1)
    assert not jm.jobs


def test_failed_job_reports_its_error():
    import concurrent.futures as cf
    import jobs
    said = []

    def run(coro):
        coro.close()
        f = cf.Future()
        f.set_exception(RuntimeError('Amazon asked for a captcha\nmore'))
        return f

    async def work(progress):
        return 1
    jm = jobs.JobManager(run, say=said.append)
    j = jm.get(jm.start('check', lambda p: work(p)))
    assert j['status'] == 'error' and 'captcha' in j['error'] and said and 'Check failed' in said[0]


def test_radar_scan_reads_the_body_once_and_answers_at_once(server, monkeypatch):
    """The market used to be read from the body twice: the second read waited for the socket to time out, then the
    scan failed with market None ("Starting…" for minutes)."""
    import asyncio
    import time as _t
    got = []
    monkeypatch.setattr(server, 'new_job', lambda kind, fn: got.append(fn) or 'job1')
    t0 = _t.time()
    st, _, raw = call(server, 'POST', '/api/radar/scan', {'market': 'AU'})
    assert st == 200 and json.loads(raw)['job'] == 'job1' and _t.time() - t0 < 2
    lines = []

    async def lists(task, progress=None):
        lines.append(task['market'])
        raise RuntimeError('stop here')
    monkeypatch.setattr(server.collectors, 'lists', lists)
    with pytest.raises(RuntimeError, match='stop here'):
        asyncio.run(got[0](lines.append))
    assert 'Amazon Australia' in lines[0] and lines[1] == 'AU'


def test_state_has_exchange_rates_and_your_ranges(server):
    st, _, raw = call(server, 'GET', '/api/state')
    fx = json.loads(raw)['fx']
    assert st == 200 and fx['currency'] == 'USD' and fx['ranges']['US'] == [20, 50]
    assert fx['ranges']['AE'][0] > 70 and fx['per_usd']['AED'] == 3.6725
    st, _, raw = call(server, 'POST', '/api/settings', {'price_currency': 'EUR', 'price_min': 10})
    d = json.loads(raw)
    assert st == 200 and d['settings']['price_min'] == 10 and d['fx']['ranges']['US'][0] > 10
    call(server, 'POST', '/api/settings', {'price_currency': 'USD', 'price_min': 20})


def test_opportunity_map_route(server):
    st, _, raw = call(server, 'GET', '/api/opportunity_map?market=AE')
    d = json.loads(raw)
    assert st == 200 and d['market'] == 'AE' and isinstance(d['points'], list) and d['local_line'] == 4.5


def test_profit_takes_the_supplier_cost_in_any_currency(server):
    import fxrates
    base = {'market': 'AU', 'price': 49.99, 'l': 30, 'w': 20, 'h': 10, 'kg': 0.5}
    st, _, raw = call(server, 'POST', '/api/profit', {**base, 'cogs_usd': 5})
    usd = json.loads(raw)
    aed = fxrates.convert(5, 'USD', 'AED')
    st2, _, raw = call(server, 'POST', '/api/profit', {**base, 'cogs': aed, 'cogs_currency': 'AED'})
    d = json.loads(raw)
    assert st == st2 == 200 and d['economics']['COGS'] == pytest.approx(usd['economics']['COGS'], abs=0.01)
    assert d['cost_currency'] == 'AED'
    assert d['max_cogs_20_cost'] == pytest.approx(fxrates.convert(usd['max_cogs_20'], 'USD', 'AED'), abs=0.02)
    st, _, raw = call(server, 'POST', '/api/profit', {**base, 'cogs': 5, 'cogs_currency': 'XYZ'})
    assert json.loads(raw)['cost_currency'] == 'USD'                       # unknown currency: read as US$


def test_check_export_has_prices_in_your_currency(server):
    import io
    import time

    import pandas as pd

    import amazon_check as ac
    import fxrates
    rows = [{'market': 'AU', 'asin': 'B%d' % i, 'title': 't', 'price': p, 'sponsored': False, 'days': 2, 'delivery': 'x',
             'reviews': 1, 'bought': None} for i, p in enumerate((10.0, 30.0))]
    s = ac.summarize('AU', 'moving bags', 'Sydney 2000', rows)
    server.jobs.jobs['xjob'] = {'id': 'xjob', 'status': 'done', 'result': {'keyword': 'moving bags', 'results': [s]},
                                'finished': time.time(), 'log': []}
    st, _, raw = call(server, 'POST', '/api/export', {'kind': 'check', 'job': 'xjob'})
    sheets = pd.read_excel(io.BytesIO(raw), sheet_name=None)
    summ, listings = sheets['Summary'], sheets['Listings']
    assert st == 200 and summ.loc[0, 'currency'] == 'A$' and summ.loc[0, 'typical_price'] == 20
    assert summ.loc[0, 'typical_price_USD'] == pytest.approx(fxrates.convert(20, 'AUD', 'USD'), abs=0.01)
    assert list(listings['currency']) == ['A$', 'A$']
    assert listings.loc[1, 'price_USD'] == pytest.approx(fxrates.convert(30, 'AUD', 'USD'), abs=0.01)


def test_list_export_checks_every_row_and_converts_prices(server):
    import io

    import openpyxl
    import pandas as pd

    import fxrates
    rows = [{'product': '=HYPERLINK("http://x")', 'market': 'AU', 'local': 3, 'overseas': 9, 'fast': 2, 'total': 20,
             'demand': 1200.5, 'per_local_seller': 300, 'overseas_share': 0.42, 'price_med': 20, 'verdict': 'gap'},
            {'product': 'dog bed', 'market': 'AE', 'error': 'captcha'},
            {'product': 'bad', 'market': 'XX', 'local': 1}, 'not a row', {'product': 'odd', 'market': 'US', 'local': 'many'}]
    st, _, raw = call(server, 'POST', '/api/export', {'kind': 'batch', 'rows': rows})
    df = pd.read_excel(io.BytesIO(raw))
    assert st == 200 and list(df['Market']) == ['Australia', 'UAE', 'USA']                  # the unknown market is left out
    assert df.loc[0, 'Overseas share of sales %'] == 42 and df.loc[1, 'Verdict'] == 'captcha'
    assert df.loc[0, 'Typical price (US$)'] == pytest.approx(fxrates.convert(20, 'AUD', 'USD'), abs=0.01)
    assert pd.isna(df.loc[2, 'Local sellers'])                                              # not a number: left empty
    ws = openpyxl.load_workbook(io.BytesIO(raw)).active
    assert ws['A2'].data_type == 's' and ws['A2'].value.startswith("'=")                     # text, never a formula
    assert call(server, 'POST', '/api/export', {'kind': 'batch', 'rows': []})[0] == 400


def test_download_name_works_for_an_arabic_keyword(server):
    import time

    import amazon_check as ac
    rows = [{'market': 'AE', 'asin': 'B1', 'title': 't', 'price': 50.0, 'sponsored': False, 'days': 2, 'delivery': 'x',
             'reviews': 1, 'bought': None}]
    kw = 'حقيبة "سفر"'
    server.jobs.jobs['ajob'] = {'id': 'ajob', 'status': 'done', 'finished': time.time(), 'log': [],
                                'result': {'keyword': kw, 'results': [ac.summarize('AE', kw, 'Dubai', rows)]}}
    c = http.client.HTTPConnection('127.0.0.1', server.port, timeout=5)
    c.request('POST', '/api/export', body=json.dumps({'kind': 'check', 'job': 'ajob'}).encode(),
              headers={'Host': '127.0.0.1:%d' % server.port, 'X-PC-Token': server.token, 'Content-Type': 'application/json'})
    r = c.getresponse()
    r.read()
    c.close()
    cd = r.getheader('Content-Disposition')
    assert r.status == 200 and 'filename="Delivery_check_.xlsx"' in cd and "filename*=UTF-8''Delivery_check_%D8" in cd
    assert app.attachment('Top_products.xlsx') == 'attachment; filename="Top_products.xlsx"; filename*=UTF-8\'\'Top_products.xlsx'


def test_sizes_are_read_from_the_top_product_pages(server, monkeypatch):
    import asyncio

    import amazon_check as ac
    import product
    import storage
    rows = [{'market': 'AU', 'asin': 'S%d' % i, 'title': 'Box %d' % i, 'price': 30.0, 'sponsored': i == 1, 'days': 2,
             'delivery': 'x', 'reviews': 5, 'bought': None} for i in range(7)]
    server.db.write(storage.save_serp, ac.summarize('AU', 'storage box', 'Sydney 2000', rows), 'user', None, 1, '2000')
    read = []

    async def fake_read(checker, asin, code):
        read.append(asin)
        if asin == 'S2':
            raise RuntimeError('page did not load')
        return {'title': 'Box', 'root_rank': 100, 'status': 'ok', 'pkg_cm': [40.0, 30.0, 10.0 + len(read)], 'pkg_kg': 1.2}
    monkeypatch.setattr(product, 'read_product', fake_read)
    sz = asyncio.run(server.read_sizes('AU', 'storage box'))
    assert read == ['S0', 'S2', 'S3'] and sz['n'] == 2 and sz['dims_cm'] == [40.0, 30.0, 12.0] and sz['kg'] == 1.2
    assert sz['missing'] == ['S2', 'S4', 'S5']                                  # the ad (S1) is never opened

    async def captcha(checker, asin, code):
        read.append(asin)
        raise ac.CaptchaRequired('Amazon wants a captcha')
    monkeypatch.setattr(product, 'read_product', captcha)
    del read[:]
    asyncio.run(server.read_sizes('AU', 'storage box', want=5))
    assert read == ['S2']                                                       # stops at the captcha
    with pytest.raises(ac.CaptchaRequired):                                     # the button says why nothing was read
        asyncio.run(server.read_sizes('AU', 'storage box', want=5, strict=True))
    st, _, raw = call(server, 'POST', '/api/parse_size', {'text': 'Single package size: 30X20X10 cm Single gross weight: 0.5 kg'})
    assert st == 200 and json.loads(raw) == {'dims_cm': [30, 20, 10], 'kg': 0.5}


def test_negative_reviews_are_read_from_amazon(server, monkeypatch):
    import asyncio

    import amazon_check as ac
    import product
    import storage
    rows = [{'market': 'US', 'asin': 'M%d' % i, 'title': 'Bowl %d' % i, 'price': 20.0, 'sponsored': i == 1, 'days': 2,
             'delivery': 'x', 'reviews': 5, 'bought': None} for i in range(4)]
    server.db.write(storage.save_serp, ac.summarize('US', 'mixing bowl', 'New York 10001', rows), 'user', None, 1, '10001')
    calls, signed = [], [False]

    async def critical(checker, asin, code):
        calls.append(('reviews page', asin))
        if not signed[0]:
            raise product.SignInNeeded('sign in')
        return [{'id': 'RX' + asin + '000000', 'stars': 1.0, 'title': 'Broke', 'body': 'b', 'local': True, 'helpful': 9}]

    async def page(checker, asin, code, reviews=False):
        calls.append(('product page', asin))
        return {'title': 'Bowl', 'root_rank': 10, 'status': 'ok', 'reviews': 50, 'signed_in': signed[0], 'review_items': [
            {'id': 'R' + asin + '00000000', 'stars': 2.0, 'title': 'Cracked', 'body': 'b', 'local': True, 'helpful': 3}]}
    monkeypatch.setattr(product, 'read_critical_reviews', critical)
    monkeypatch.setattr(product, 'read_product', page)
    st, _, raw = call(server, 'GET', '/api/report/reviews?market=US&kw=mixing%20bowl')
    assert st == 200 and json.loads(raw)['todo'] == ['M0', 'M2', 'M3'] and json.loads(raw)['items'] == []
    assert asyncio.run(server.read_reviews('US', ['M0', 'M2', 'M3'])) == 3
    assert calls == [('reviews page', 'M0'), ('product page', 'M0'), ('product page', 'M2'), ('product page', 'M3')]
    r = json.loads(call(server, 'GET', '/api/report/reviews?market=US&kw=mixing%20bowl')[2])
    assert r['todo'] == [] and [i['asin'] for i in r['items']] == ['M0', 'M2', 'M3'] and not r['critical']
    signed[0], calls[:] = True, []                          # signed in now: a product page shows it, then reviews pages
    asyncio.run(server.read_reviews('US', ['M0', 'M2']))
    assert calls == [('product page', 'M0'), ('reviews page', 'M2')]
    r = json.loads(call(server, 'GET', '/api/report/reviews?market=US&asin=M2')[2])
    assert r['critical'] and r['items'][0]['title'] == 'Broke'

    async def captcha(checker, asin, code, reviews=False):
        raise ac.CaptchaRequired('Amazon wants a captcha')
    monkeypatch.setattr(product, 'read_critical_reviews', captcha)
    with pytest.raises(ac.CaptchaRequired):                                     # nothing read: the reason is shown
        asyncio.run(server.read_reviews('US', ['M3']))
    assert call(server, 'POST', '/api/report/reviews', {'market': 'US', 'asin': 'B0-1'})[0] == 400
    assert call(server, 'GET', '/api/report/reviews?market=US&kw=%20')[0] == 400
